"""XLSX parser. Wraps `openpyxl` only.

CSV/TSV lives in the companion module `parse/csv.py` (stdlib `csv`).

Format key: `"xlsx"`.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from okg.substrate.library.extraction._protocols import (
    ParsedChunk,
    ParsedDocument,
    _strip_control_chars,
)
from okg.substrate.library.extraction.parse.dispatch import register_parser


class XlsxParser:
    """`.xlsx` -> ParsedDocument via openpyxl."""

    def parse(self, content: bytes | str, **_: Any) -> ParsedDocument:
        if isinstance(content, str):
            content = content.encode("utf-8")
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
        sheets = wb.sheetnames
        # `text` is every sheet's rows joined by "\n"; each sheet chunk
        # carries its span in it (an empty sheet adds no row, so its span is
        # empty and no separator is counted for it).
        chunks: list[ParsedChunk] = []
        all_rows: list[str] = []
        position = 0
        for sheet_name in sheets:
            ws = wb[sheet_name]
            sheet_rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                sheet_rows.append(_strip_control_chars(
                    "\t".join("" if v is None else str(v) for v in row)
                ))
            sheet_text = "\n".join(sheet_rows)
            if sheet_rows and all_rows:
                position += 1
            chunks.append(ParsedChunk(
                text=sheet_text,
                heading_path=[sheet_name],
                span=(position, position + len(sheet_text)),
            ))
            all_rows.extend(sheet_rows)
            position += len(sheet_text)
        return ParsedDocument(
            text="\n".join(all_rows),
            metadata={
                "sheets": sheets,
                "sheet_count": len(sheets),
            },
            chunks=chunks,
        )


register_parser("xlsx", XlsxParser)


__all__ = ["XlsxParser"]
