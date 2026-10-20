"""`.docx` parser. Wraps `python-docx`."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from okg.substrate.library.extraction._protocols import (
    ParsedChunk,
    ParsedDocument,
    _strip_control_chars,
)
from okg.substrate.library.extraction.parse.dispatch import register_parser


class DocxParser:
    """`.docx` -> ParsedDocument via python-docx."""

    def parse(self, content: bytes | str, **_: Any) -> ParsedDocument:
        if isinstance(content, str):
            content = content.encode("utf-8")
        import docx as _docx

        doc = _docx.Document(BytesIO(content))
        # `text` is every paragraph, headings included, joined by "\n"; each
        # body chunk carries its span in it. Headings are not chunks, so a
        # whole read shows them as gaps and says it is partial rather than
        # complete.
        chunks: list[ParsedChunk] = []
        text_parts: list[str] = []
        heading_path: list[str] = []
        position = 0
        for para in doc.paragraphs:
            txt = _strip_control_chars((para.text or "").strip())
            if not txt:
                continue
            if text_parts:
                position += 1
            style = (para.style.name if para.style else "") or ""
            if style.startswith("Heading "):
                try:
                    level = int(style.split()[-1])
                except (ValueError, IndexError):
                    level = 1
                heading_path = heading_path[: level - 1]
                heading_path.append(txt)
            else:
                chunks.append(ParsedChunk(
                    text=txt,
                    heading_path=list(heading_path),
                    span=(position, position + len(txt)),
                ))
            text_parts.append(txt)
            position += len(txt)

        meta = doc.core_properties
        metadata: dict[str, Any] = {
            "title": meta.title or None,
            "author": meta.author or None,
            "created": meta.created.isoformat() if meta.created else None,
        }
        metadata = {k: v for k, v in metadata.items() if v is not None}

        return ParsedDocument(
            text="\n".join(text_parts),
            metadata=metadata,
            chunks=chunks,
        )


register_parser("docx", DocxParser)


__all__ = ["DocxParser"]
